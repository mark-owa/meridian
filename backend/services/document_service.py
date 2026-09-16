"""Document upload, storage, and text extraction.

Upload pipeline: validate -> stream to disk -> create DB record.
Text extraction (step 2, run in the background worker) differs per format:
PDFs are read page-by-page, DOCX by paragraph, plain text directly.
"""

import uuid
from pathlib import Path

from fastapi import HTTPException, UploadFile, status
from sqlalchemy.orm import Session

from config import get_settings
from models.database import Document
from utils.security import validate_file_type
from utils.time import utcnow

settings = get_settings()

ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "text/plain": "txt",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/markdown": "txt",
}

UPLOAD_DIR = Path(settings.UPLOAD_DIR)
UPLOAD_DIR.mkdir(exist_ok=True, parents=True)


async def save_upload(file: UploadFile, owner_id: str, db: Session) -> Document:
    """Validate, stream to disk (checking size as bytes arrive so we never
    buffer an oversized file in memory), and create the DB record."""
    content_type = file.content_type or ""
    if content_type not in ALLOWED_TYPES:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"File type '{content_type}' not supported. Use PDF, TXT, or DOCX.",
        )

    declared_extension = ALLOWED_TYPES[content_type]
    safe_filename = f"{uuid.uuid4()}.{declared_extension}"
    file_path = UPLOAD_DIR / safe_filename
    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024

    total_bytes = 0
    try:
        with open(file_path, "wb") as f:
            while chunk := await file.read(8192):
                total_bytes += len(chunk)
                if total_bytes > max_bytes:
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=f"File exceeds {settings.MAX_UPLOAD_SIZE_MB}MB limit",
                    )
                f.write(chunk)
    except BaseException:
        # Include interrupted uploads; never keep a partially written file.
        file_path.unlink(missing_ok=True)
        raise

    # The client-supplied Content-Type header is never sufficient. Validate
    # the actual bytes even when optional libmagic support is unavailable.
    if not validate_file_type(file_path, declared_extension):
        file_path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="File content doesn't match its declared type",
        )

    doc = Document(
        owner_id=owner_id,
        filename=safe_filename,
        original_name=file.filename or "unnamed",
        file_type=declared_extension,
        file_size_bytes=total_bytes,
        status="pending",
    )
    try:
        db.add(doc)
        db.commit()
    except Exception:
        db.rollback()
        file_path.unlink(missing_ok=True)
        raise
    db.refresh(doc)
    return doc


def extract_text_from_file(file_path: Path, file_type: str) -> str:
    if file_type == "pdf":
        return _extract_pdf(file_path)
    elif file_type == "docx":
        return _extract_docx(file_path)
    elif file_type == "txt":
        return _extract_txt(file_path)
    raise ValueError(f"Unknown file type: {file_type}")


def _extract_pdf(file_path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    pages = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text()
        if text and text.strip():
            pages.append(f"[Page {page_num}]\n{text.strip()}")

    if not pages:
        raise ValueError("PDF contains no extractable text (it may be a scanned image)")
    return "\n\n".join(pages)


def _extract_docx(file_path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _extract_txt(file_path: Path) -> str:
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="latin-1")


def get_document(db: Session, doc_id: str, owner_id: str) -> Document:
    doc = db.query(Document).filter(Document.id == doc_id, Document.owner_id == owner_id).first()
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Document {doc_id} not found or access denied",
        )
    return doc


def list_documents(db: Session, owner_id: str, skip: int = 0, limit: int = 50) -> tuple[list, int]:
    query = db.query(Document).filter(Document.owner_id == owner_id)
    total = query.count()
    docs = query.order_by(Document.created_at.desc()).offset(skip).limit(limit).all()
    return docs, total


def delete_document(db: Session, doc_id: str, owner_id: str) -> None:
    doc = get_document(db, doc_id, owner_id)
    file_path = UPLOAD_DIR / doc.filename
    if file_path.exists():
        file_path.unlink()
    db.delete(doc)
    db.commit()


def update_document_status(
    db: Session,
    doc_id: str,
    status: str,
    chunk_count: int = 0,
    error_message: str | None = None,
) -> Document | None:
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if doc:
        doc.status = status
        doc.chunk_count = chunk_count
        doc.error_message = error_message
        doc.processed_at = utcnow() if status in ("ready", "failed") else None
        db.commit()
        db.refresh(doc)
    return doc
